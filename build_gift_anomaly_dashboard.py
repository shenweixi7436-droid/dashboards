# -*- coding: utf-8 -*-
"""从 费用分析-样品及活动.xlsx 生成 GIFT_LIVE_DATA 数据文件，供独立看板顶部"实时数据"横幅使用。

数据源切换策略：
- 主看板「线下赠品稽核」小卡片继续读旧源（市场稽核部重点工作.xlsx / 赠品稽核-活动、赠品稽核-样品）
- 独立看板（线下活动与样品申请异常解读看板.html）顶部新增的实时数据横幅读新源（本脚本）
- 原有的"异常解读"模块基于内嵌的 anomalyData 快照保留

输出（两个位置同时写，保持同步）：
- 仓库部署位：github-dashboards/assets/pages/gift-live-data.js —— 主看板点击卡片跳转到该页，须与页面同目录，且只能走仓库内相对路径（file:// 绝对路径在 http/GitHub Pages 下会被浏览器拦截）
- 本地工作副本：<独立看板所在目录>/gift-live-data.js
- 变量：window.GIFT_LIVE_DATA
"""
from __future__ import annotations

import json
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

SOURCE = Path(r"C:/Users/shenw/Desktop/看板/费用分析-样品及活动.xlsx")
OUTPUT = Path(r"C:/Users/shenw/WorkBuddy/2026-08-25-13-53-45/gift-live-data.js")
REPO_OUTPUT = Path(
    r"C:/Users/shenw/Documents/New project/github-dashboards/assets/pages/gift-live-data.js"
)

ACT_SHEET = "线下活动试吃品申请单"
SAM_SHEET = "线下样品申请单"

# 活动：col 30 = 商品明细:预估金额(元)（每行商品金额=单价×数量，可直接累加）；
#       col 31 = 申请单维度总金额(冗余重复填在每行，不能累加)；
#       col 25 = 申请使用场景
ACT_AMOUNT_COL = 30
ACT_USAGE_COL = 25
# 样品：col 27 = 商品明细:预估金额(元)（每行商品金额，可直接累加）；col 20 = 样品用途
SAM_AMOUNT_COL = 27
SAM_USAGE_COL = 20


def num(v):
    if v in (None, ""):
        return 0.0
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def month_key(v):
    if v in (None, ""):
        return None
    if isinstance(v, datetime):
        return v.month
    try:
        m = int(float(v))
        if 1 <= m <= 12:
            return m
    except (TypeError, ValueError):
        pass
    return None


def aggregate(ws, amount_col, usage_col, label_map):
    """按月份 + 申请使用场景/样品用途 汇总。

    label_map: 源表分类名 -> 归一化展示名（如 {"经销商搭赠":"经销商搭赠", ...}）
    """
    by_month_total = {}  # month -> {count, amount}
    by_month_category = {}  # month -> {category -> {count, amount}}
    skipped = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row:
            continue
        m = month_key(row[0])
        if m is None:
            skipped += 1
            continue
        amt = num(row[amount_col - 1] if len(row) >= amount_col else 0)
        usage_raw = row[usage_col - 1] if len(row) >= usage_col else ""
        usage = (str(usage_raw).strip() if usage_raw is not None else "")
        category = label_map.get(usage, "其他")

        bm = by_month_total.setdefault(m, {"count": 0, "amount": 0.0})
        bm["count"] += 1
        bm["amount"] += amt

        bc = by_month_category.setdefault(m, {}).setdefault(
            category, {"count": 0, "amount": 0.0}
        )
        bc["count"] += 1
        bc["amount"] += amt
    return by_month_total, by_month_category, skipped


def latest_month(by_month):
    return max(by_month.keys()) if by_month else None


def mom_pct(current, previous):
    if not previous or previous == 0:
        return None
    return round((current - previous) / previous * 100, 1)


def main():
    if not SOURCE.exists():
        print(f"ERROR: source not found: {SOURCE}")
        sys.exit(1)
    wb = load_workbook(SOURCE, data_only=True, read_only=True)

    # 活动 4 类：经销商搭赠 / 消费者搭赠 / 消费者试吃 / 其他
    act_label_map = {
        "经销商搭赠": "经销商搭赠",
        "消费者搭赠": "消费者搭赠",
        "消费者试吃": "消费者试吃",
    }
    act_total, act_cat, act_skip = aggregate(
        wb[ACT_SHEET], ACT_AMOUNT_COL, ACT_USAGE_COL, act_label_map
    )

    # 样品 2 类：老客户增加品类送样 / 新客户招商送样
    sam_label_map = {
        "老客户增加品类送样": "老客户增加品类送样",
        "新客户招商送样": "新客户招商送样",
    }
    sam_total, sam_cat, sam_skip = aggregate(
        wb[SAM_SHEET], SAM_AMOUNT_COL, SAM_USAGE_COL, sam_label_map
    )

    # 推导 T-1 月 = 详情表最新月
    act_t1 = latest_month(act_total)
    sam_t1 = latest_month(sam_total)
    if act_t1 != sam_t1:
        print(f"WARN: 活动 T-1月({act_t1}) 与 样品 T-1月({sam_t1}) 不一致")
    t1 = act_t1 or sam_t1

    act_t1_data = act_total.get(t1, {"count": 0, "amount": 0.0})
    sam_t1_data = sam_total.get(t1, {"count": 0, "amount": 0.0})
    act_t0_data = act_total.get(t1 - 1, {"count": 0, "amount": 0.0}) if t1 and t1 > 1 else None
    sam_t0_data = sam_total.get(t1 - 1, {"count": 0, "amount": 0.0}) if t1 and t1 > 1 else None

    payload = {
        "source": "费用分析-样品及活动.xlsx",
        "generatedAt": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "t1Month": f"{t1}月" if t1 else None,
        "activity": {
            "monthly": {f"{m}月": v for m, v in sorted(act_total.items())},
            "byCategoryMonthly": {
                f"{m}月": v for m, v in sorted(act_cat.items())
            },
            "t1": act_t1_data,
            "t0": act_t0_data,
            "countMomPct": mom_pct(act_t1_data["count"], act_t0_data["count"]) if act_t0_data else None,
            "amountMomPct": mom_pct(act_t1_data["amount"], act_t0_data["amount"]) if act_t0_data else None,
        },
        "sample": {
            "monthly": {f"{m}月": v for m, v in sorted(sam_total.items())},
            "byCategoryMonthly": {
                f"{m}月": v for m, v in sorted(sam_cat.items())
            },
            "t1": sam_t1_data,
            "t0": sam_t0_data,
            "countMomPct": mom_pct(sam_t1_data["count"], sam_t0_data["count"]) if sam_t0_data else None,
            "amountMomPct": mom_pct(sam_t1_data["amount"], sam_t0_data["amount"]) if sam_t0_data else None,
        },
        "skippedRows": {"activity": act_skip, "sample": sam_skip},
    }

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    REPO_OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    js_body = (
        f"// 自动生成自 build_gift_anomaly_dashboard.py（{payload['generatedAt']}）\n"
        f"window.GIFT_LIVE_DATA = {json.dumps(payload, ensure_ascii=False, indent=2)};\n"
    )
    OUTPUT.write_text(js_body, encoding="utf-8")
    REPO_OUTPUT.write_text(js_body, encoding="utf-8")
    print(
        f"GIFT_LIVE_DATA written: t1={payload['t1Month']}, "
        f"活动 {act_t1_data['count']}条/{act_t1_data['amount']:.0f}元, "
        f"样品 {sam_t1_data['count']}条/{sam_t1_data['amount']:.0f}元\n"
        f"  out1(工作副本)={OUTPUT}\n"
        f"  out2(仓库部署)={REPO_OUTPUT}"
    )


if __name__ == "__main__":
    main()
