from collections import defaultdict
from pathlib import Path
import json

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent
SOURCE = Path(r"C:\Users\shenw\Desktop\看板\市场稽核部重点工作.xlsx")
OUT = ROOT / "assets" / "data" / "device-region-analysis.js"
DETAIL_SHEET = "智能设备明细"
EXCLUDED_SYSTEMS = {"零食", "总部"}


def text(value):
    return "" if value is None else str(value).strip()


def device_kind(value):
    name = text(value)
    if "保温柜" in name:
        return "warm"
    if "烤肠机" in name:
        return "sausage"
    return ""


def find_header_row(ws):
    for row in range(1, min(ws.max_row, 30) + 1):
        values = [text(ws.cell(row, col).value) for col in range(1, min(ws.max_column, 30) + 1)]
        if "是否开机" in values and any("达标" in value for value in values):
            return row, values
    raise ValueError("未找到同时包含“是否开机”和“是否达标”的表头行")


def aggregate(rows, key_name, top=None):
    buckets = defaultdict(lambda: {
        "warmOnOrVolume": 0,
        "warmBad": 0,
        "sausageOnOrVolume": 0,
        "sausageBad": 0,
    })
    for row in rows:
        key = row[key_name] or "未填写"
        bucket = buckets[key]
        prefix = "warm" if row["device"] == "warm" else "sausage"
        bucket[prefix + "OnOrVolume"] += 1
        if row["bad"]:
            bucket[prefix + "Bad"] += 1
    result = []
    for dimension, values in buckets.items():
        base = values["warmOnOrVolume"] + values["sausageOnOrVolume"]
        bad = values["warmBad"] + values["sausageBad"]
        result.append({"dimension": dimension, **values, "rate": round(bad / base * 100, 1) if base else 0})
    result.sort(key=lambda item: (-item["rate"], -item["warmBad"] - item["sausageBad"], item["dimension"]))
    return result[:top] if top else result


def build():
    wb = load_workbook(SOURCE, data_only=True, read_only=True)
    if DETAIL_SHEET not in wb.sheetnames:
        return {
            "available": False,
            "message": "源文件中未找到“智能设备明细”工作表。当前仅有“智能设备台账汇总”，无法按省份和客户生成明细分析。",
            "source": SOURCE.name,
        }
    ws = wb[DETAIL_SHEET]
    header_row, headers = find_header_row(ws)
    qualified_col = next((index + 1 for index, value in enumerate(headers) if "是否达标" in value), None)
    if not qualified_col:
        raise ValueError("未找到“是否达标”列")

    base_rows = []
    qualified_rows = []
    for values in ws.iter_rows(min_row=header_row + 1, values_only=True):
        province = text(values[0] if len(values) > 0 else "")
        system = text(values[1] if len(values) > 1 else "")
        customer = text(values[2] if len(values) > 2 else "")
        is_on = text(values[4] if len(values) > 4 else "")
        device = device_kind(values[9] if len(values) > 9 else "")
        qualified = text(values[qualified_col - 1] if len(values) >= qualified_col else "")
        if system in EXCLUDED_SYSTEMS or not device or (not province and not customer):
            continue
        common = {"province": province, "customer": customer, "device": device}
        base_rows.append({**common, "bad": is_on == "否"})
        if is_on == "是":
            qualified_rows.append({**common, "bad": qualified == "否"})

    return {
        "available": True,
        "source": f"{SOURCE.name} / {DETAIL_SHEET}",
        "provinceNotOn": aggregate(base_rows, "province"),
        "customerNotOn": aggregate(base_rows, "customer", 20),
        "provinceNotQualified": aggregate(qualified_rows, "province"),
        "customerNotQualified": aggregate(qualified_rows, "customer", 20),
    }


def main():
    payload = build()
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text("window.DEVICE_REGION_ANALYSIS = " + json.dumps(payload, ensure_ascii=False, indent=2) + ";\n", encoding="utf-8")
    print(f"已生成: {OUT}")
    print(payload.get("message") or "智能设备区域分析数据生成完成")


if __name__ == "__main__":
    main()
